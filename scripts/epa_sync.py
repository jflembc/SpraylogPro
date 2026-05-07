#!/usr/bin/env python3
"""
EPA APPRIL public dump -> Supabase chemical_library
Uses openpyxl read_only mode to avoid Node/JSZip string size limits on large .xlsx files.

Env:
  SUPABASE_URL
  SUPABASE_SERVICE_ROLE_KEY
  MANUAL_RUN=1  optional; bypass Monday 7PM ET gate + weekly success check (manual workflow)
"""

from __future__ import annotations

import io
import os
import sys
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

import requests
from openpyxl import load_workbook
from supabase import create_client

EPA_DUMP_URL = "https://www3.epa.gov/pesticides/appril/apprildatadump_public.xlsx"
SHEET_PREFERRED = "ACTIVE_REGISTRATIONS5_PUBLIC_V"

SUPABASE_URL = os.environ.get("SUPABASE_URL", "").strip()
SUPABASE_SERVICE_ROLE_KEY = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
MANUAL_RUN = os.environ.get("MANUAL_RUN", "").strip() == "1"


def die(msg: str, code: int = 1) -> None:
    print(msg, file=sys.stderr)
    sys.exit(code)


if not SUPABASE_URL or not SUPABASE_SERVICE_ROLE_KEY:
    die("Missing SUPABASE_URL or SUPABASE_SERVICE_ROLE_KEY")


def norm_key(s: str) -> str:
    return "".join(c for c in s.lower() if c.isalnum())


def get_loose(rec: Dict[str, Any], *keys: str) -> Any:
    idx = {norm_key(k): v for k, v in rec.items()}
    for k in keys:
        nk = norm_key(k)
        if nk in idx and idx[nk] is not None and str(idx[nk]).strip() != "":
            return idx[nk]
    return None


def iso_week_key(dt: Optional[datetime] = None) -> str:
    dt = dt or datetime.now(timezone.utc)
    # ISO week (simple, good enough for a "week bucket" label)
    y, w, _ = dt.isocalendar()
    return f"{y}-W{w:02d}"


def ny_monday_7pm_open() -> bool:
    from zoneinfo import ZoneInfo

    z = ZoneInfo("America/New_York")
    now = datetime.now(z)
    return now.weekday() == 0 and now.hour == 19  # Monday 19:00


def cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (str, int, float, bool)):
        return str(v).strip()
    return str(v).strip()


def map_row(r: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    reg = cell_str(get_loose(r, "reg_num", "registration number", "registration_number"))
    name = cell_str(get_loose(r, "product_name", "product name"))
    if not reg or not name:
        return None

    ai = cell_str(
        get_loose(
            r,
            "ais",
            "active ingredient(s)",
            "active ingredients",
            "active_ingredient",
        )
    )

    return {
        "reg_num": reg,
        "epa_number": reg,
        "product_name": name,
        "active_ingredient": ai,
        "active_ingredients": ai,
        "signal_word": cell_str(get_loose(r, "signal_word", "signal word")),
        "product_type": cell_str(
            get_loose(r, "pesticide_type", "pesticide type", "product_type", "reg_type", "reg type")
        ),
        "pesticide_type": cell_str(get_loose(r, "pesticide_type", "pesticide type")),
        "status_group": cell_str(get_loose(r, "status_group", "status group")),
        "status": cell_str(get_loose(r, "status")),
        "use_pattern": cell_str(get_loose(r, "use_pattern", "use pattern")),
        "company_name": cell_str(get_loose(r, "company_name", "company name")),
        "source": "EPA_APPRIL",
        "last_synced_at": datetime.now(timezone.utc).isoformat(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


def iter_dump_rows(path: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET_PREFERRED in wb.sheetnames:
            ws = wb[SHEET_PREFERRED]
        else:
            ws = wb[wb.sheetnames[0]]

        it = ws.iter_rows(values_only=True)
        try:
            header_row = next(it)
        except StopIteration:
            return

        headers = [cell_str(h) for h in header_row if cell_str(h)]
        col_count = len(header_row)

        for row in it:
            obj: Dict[str, Any] = {}
            for i, h in enumerate(header_row):
                if i >= col_count:
                    break
                hn = cell_str(h)
                if not hn:
                    continue
                val = row[i] if i < len(row) else None
                obj[hn] = val
            yield obj
    finally:
        wb.close()


def main() -> None:
    sb = create_client(SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)
    week_key = iso_week_key()

    if not MANUAL_RUN:
        if not ny_monday_7pm_open():
            print("Skipping: not Monday 7:00 PM America/New_York.")
            return

        done = (
            sb.table("chemical_sync_runs")
            .select("id")
            .eq("week_key", week_key)
            .eq("status", "success")
            .limit(1)
            .execute()
        )
        if done.data:
            print("Already synced this week:", week_key)
            return
    else:
        print("MANUAL_RUN=1 -> bypass time gate and weekly duplicate check.")

    run_ins = (
        sb.table("chemical_sync_runs")
        .insert(
            {
                "run_type": "manual" if MANUAL_RUN else "scheduled",
                "week_key": week_key,
                "status": "running",
                "message": "EPA GitHub sync started (Python)",
            }
        )
        .execute()
    )
    if not run_ins.data:
        die("Failed to insert chemical_sync_runs row")
    run_id = run_ins.data[0]["id"]

    tmp_path = "/tmp/apprildatadump_public.xlsx"

    try:
        print("Fetching EPA dump...")
        r = requests.get(EPA_DUMP_URL, timeout=600, stream=True)
        r.raise_for_status()
        total = 0
        with open(tmp_path, "wb") as f:
            for chunk in r.iter_content(chunk_size=1024 * 1024):
                if chunk:
                    f.write(chunk)
                    total += len(chunk)
        print(f"Downloaded {total / 1024 / 1024:.1f} MB")

        print("Parsing with openpyxl (read_only)...")
        fetched = 0
        filtered = 0
        upserted = 0
        batch: List[Dict[str, Any]] = []
        BATCH = 500

        for rec in iter_dump_rows(tmp_path):
            fetched += 1
            s = cell_str(get_loose(rec, "status_group", "status group")).lower()
            t = cell_str(get_loose(rec, "reg_type", "reg type")).lower()
            if s != "active" or t != "sec3":
                continue
            filtered += 1
            m = map_row(rec)
            if not m:
                continue
            batch.append(m)
            if len(batch) >= BATCH:
                sb.table("chemical_library").upsert(batch, on_conflict="reg_num").execute()
                upserted += len(batch)
                batch.clear()
                if upserted % 5000 == 0:
                    print(f"Upserted {upserted}...")

        if batch:
            sb.table("chemical_library").upsert(batch, on_conflict="reg_num").execute()
            upserted += len(batch)

        sb.table("chemical_sync_runs").update(
            {
                "status": "success",
                "rows_fetched": fetched,
                "rows_upserted": upserted,
                "message": "EPA GitHub sync complete (Python)",
                "finished_at": datetime.now(timezone.utc).isoformat(),
            }
        ).eq("id", run_id).execute()

        print(f"Done. week={week_key} scanned={fetched} filtered_active_sec3={filtered} upserted={upserted}")
    except Exception as e:
        sb.table("chemical_sync_runs").update(
            {
                "status": "failed",
                "message": str(e),
                "finished_at": datetime.now(timezone.utc).isoformat(),
            }
        ).eq("id", run_id).execute()
        raise
    finally:
        try:
            os.remove(tmp_path)
        except OSError:
            pass


if __name__ == "__main__":
    main()
